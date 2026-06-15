from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Iterable, Mapping
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TIPOS_ARCHIVO = {
    "exported_logs": "exported-logs",
    "grouped_node_report": "grouped-node-report",
    "exported_request_DW": "exported-request_dw",
    "exported_drawing": "exported-drawing",
}

ARCHIVOS_OBLIGATORIOS = {
    "grouped_node_report",
    "exported_request_DW",
    "exported_drawing",
}

COLUMNAS_REQUERIDAS = {
    "exported_request_DW": {
        "request_code",
        "status",
        "process__subprocess",
        "drawing_user_last_name",
    },
    "exported_drawing": {
        "request_code",
        "status",
        "process.subprocess",
        "assigned_draftsman",
    },
    "grouped_node_report": {"request_code"},
    "exported_logs": {
        "request_code",
        "initial_state",
        "final_state",
        "date_created",
    },
}

COLUMNAS_ASIGNADOR = [
    "QR",
    "FLUJO",
    "ORIGEN",
    "ELEMENTOS",
    "ESTADO",
    "LLEGADA",
    "USER",
]

PROFESIONALES_INICIALES = [
    ("4264437", "Adrian Ignacio Sanchez Vacca", "dibujante.agp6@ocaglobal.com"),
    ("1233691570", "Daniela Forero Fierro", "daniela.forero@ocaglobal.com"),
    ("1053845070", "David Andres Luna Yela", "david.luna@ocaglobal.com"),
    ("1012459639", "Deisy Lorena Ahumada Becerra", "deisy.ahumada@ocaglobal.com"),
    ("1023959391", "Eduard Steven Ramos Tello", "eduard.ramost@ocaglobal.com"),
    (
        "1012354205",
        "Francisco Alexander Guerrero Urrego",
        "francisco.guerrero@ocaglobal.com",
    ),
    ("1023014624", "Jenny Lorena Porras Cano", "jenny.porras@ocaglobal.com"),
    ("1033708860", "Johan Stiven Silva Silva", "johan.silva@ocaglobal.com"),
    ("1022433323", "Jorge Luis Urrego Castañeda", "dibujante.agp1@ocaglobal.com"),
    (
        "79170578",
        "Jose Pablo Siachoque Rodriguez",
        "dibujante.agp8@ocaglobal.com",
    ),
    ("1014301686", "Juan Pablo Quiroga Losada", "juan.quiroga@ocaglobal.com"),
    (
        "1013670569",
        "Juan Sebastian Fernandez Buitrago",
        "juan.fernandezb@ocaglobal.com",
    ),
    (
        "1024565448",
        "Kenneth Gibelli Cristancho Cascante",
        "keneth.cristancho@ocaglobal.com",
    ),
    ("1012442896", "Kevin Alejandro Torres Perdomo", "kevin.torres@ocaglobal.com"),
    ("1030606383", "Lizeth Paola Avila Ariza", "lizeth.avila@ocaglobal.com"),
    (
        "1030670518",
        "Luis Alejandro Espinosa Patarroyo",
        "dibujante.agp10@ocaglobal.com",
    ),
    ("1000988576", "Luisa Fernanda Guatavita Perez", "luisa.guatavita@ocaglobal.com"),
    (
        "1049796769",
        "Manuel Fernando Sanchez Mondragon",
        "dibujante.agp22@ocaglobal.com",
    ),
    (
        "1077976852",
        "Stefany Colorado Fresneda",
        "dibujante.agp14@ocaglobal.com",
    ),
    (
        "1233691463",
        "Valery Steffany Cagua Marroquin",
        "valery.cagua@ocaglobal.com",
    ),
    (
        "1053829105",
        "William Fernando Uribe Bayona",
        "dibujante.agp15@ocaglobal.com",
    ),
    ("1022385145", "Yesid Merardo Reyes Tique", "dibujante.agp21@ocaglobal.com"),
    ("1019108406", "Yesika Arevalo Pinzon", "yesika.arevalo@ocaglobal.com"),
    ("1030627620", "Yiceth Loreyna Ayala Guio", "yiceth.ayala@ocaglobal.com"),
]

MAX_FILAS_DATOS_EXCEL = 1_048_575

REGLAS_LLEGADA = {
    "Flujo viejo": {
        "activos": {"2.5", "2.6"},
        "salidas": {"2", "2.6.1", "2.7"},
    },
    "Flujo nuevo": {
        "activos": {"drawing_request", "assigned_draftsman"},
        "salidas": {"verify_formats", "drawing_review"},
    },
}

ESTADOS_SEGUIDOS_LLEGADA = {
    estado
    for regla in REGLAS_LLEGADA.values()
    for estado in (*regla["activos"], *regla["salidas"])
}

RE_CODIGO_ESTADO = re.compile(r"^\s*([0-9]+(?:\.[0-9]+)*)")


class ErrorDatos(ValueError):
    """Error legible para validaciones de archivos y datos."""


def _nombre_archivo(archivo: object) -> str:
    nombre = getattr(archivo, "name", None)
    if nombre:
        return Path(str(nombre)).name
    return Path(str(archivo)).name


def identificar_archivos(
    archivos: Iterable[object],
) -> tuple[dict[str, object], dict[str, list[str]], list[str]]:
    encontrados: dict[str, object] = {}
    duplicados: dict[str, list[str]] = {}
    no_reconocidos: list[str] = []

    for archivo in archivos:
        nombre = _nombre_archivo(archivo)
        nombre_normalizado = nombre.casefold()
        coincidencias = [
            tipo
            for tipo, patron in TIPOS_ARCHIVO.items()
            if patron in nombre_normalizado
        ]

        if not coincidencias:
            no_reconocidos.append(nombre)
            continue

        tipo = coincidencias[0]
        if tipo in encontrados:
            duplicados.setdefault(tipo, [_nombre_archivo(encontrados[tipo])])
            duplicados[tipo].append(nombre)
        else:
            encontrados[tipo] = archivo

    return encontrados, duplicados, no_reconocidos


def _obtener_bytes(archivo: object) -> bytes:
    if isinstance(archivo, (str, Path)):
        return Path(archivo).read_bytes()
    if isinstance(archivo, bytes):
        return archivo
    if hasattr(archivo, "getvalue"):
        return archivo.getvalue()
    if hasattr(archivo, "read"):
        posicion = archivo.tell() if hasattr(archivo, "tell") else None
        contenido = archivo.read()
        if posicion is not None and hasattr(archivo, "seek"):
            archivo.seek(posicion)
        return contenido
    raise TypeError(f"No se puede leer el archivo de tipo {type(archivo).__name__}.")


def _detectar_codificacion(contenido: bytes) -> tuple[str, str]:
    muestra = contenido[:200_000]
    for codificacion in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            return codificacion, muestra.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return "latin1", muestra.decode("latin1")


def _detectar_separador(texto: str) -> str:
    try:
        dialecto = csv.Sniffer().sniff(texto[:10_000], delimiters=",;\t|")
        return dialecto.delimiter
    except csv.Error:
        conteos = {separador: texto[:10_000].count(separador) for separador in ",;\t|"}
        return max(conteos, key=conteos.get)


def leer_csv(archivo: object) -> pd.DataFrame:
    contenido = _obtener_bytes(archivo)
    codificacion, muestra = _detectar_codificacion(contenido)
    separador = _detectar_separador(muestra)

    try:
        df = pd.read_csv(
            io.BytesIO(contenido),
            sep=separador,
            encoding=codificacion,
            dtype=str,
            keep_default_na=False,
            low_memory=False,
            escapechar="\\",
        )
    except Exception as exc:
        raise ErrorDatos(
            f"No fue posible leer {_nombre_archivo(archivo)} como CSV. "
            f"Detalle: {exc}"
        ) from exc

    df.columns = [str(columna).strip().lstrip("\ufeff") for columna in df.columns]
    return df


def validar_columnas(
    df: pd.DataFrame, columnas_requeridas: Iterable[str], nombre_archivo: str
) -> None:
    requeridas = set(columnas_requeridas)
    faltantes = sorted(requeridas.difference(df.columns))
    if faltantes:
        disponibles = ", ".join(map(str, df.columns)) or "(ninguna)"
        raise ErrorDatos(
            f"El archivo {nombre_archivo} no contiene: {', '.join(faltantes)}. "
            f"Columnas disponibles: {disponibles}."
        )


def _texto_sin_tildes(valor: object) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(caracter for caracter in texto if not unicodedata.combining(caracter))


def normalizar_origen(valor: object) -> str:
    texto = "" if valor is None else str(valor).strip()
    comparable = _texto_sin_tildes(texto).casefold()
    if any(patron in comparable for patron in ("mayor impacto", "averias bt")):
        return "E1"
    return texto


def formatear_nombre_flujo_viejo(valor: object) -> str:
    texto = "" if valor is None else re.sub(r"\s+", " ", str(valor)).strip()
    if not texto or "," not in texto:
        return ""

    apellidos, nombres = texto.split(",", maxsplit=1)
    apellidos = apellidos.strip(" ,")
    nombres = nombres.strip(" ,")
    if not apellidos or not nombres:
        return ""
    return f"{nombres} {apellidos}".title()


def limpiar_cedula(valor: object) -> str:
    if valor is None or pd.isna(valor):
        return ""
    texto = str(valor).strip()
    if re.fullmatch(r"\d+\.0", texto):
        texto = texto[:-2]
    return "".join(re.findall(r"\d", texto))


def es_estado_flujo_viejo(valor: object) -> bool:
    texto = "" if valor is None else str(valor).strip()
    return bool(re.match(r"^2\.(?:5|6)(?:\D|$)", texto))


def normalizar_estado_flujo_nuevo(valor: object) -> str:
    texto = "" if valor is None else str(valor).strip()
    comparable = _texto_sin_tildes(texto).casefold().replace(" ", "_")
    if comparable in {"assigned_draftsman", "drawing_request"}:
        return comparable
    return ""


def profesionales_por_defecto() -> pd.DataFrame:
    return pd.DataFrame(
        PROFESIONALES_INICIALES,
        columns=["CEDULA", "NOMBRE", "CORREO"],
        dtype=str,
    )


def normalizar_profesionales(df: pd.DataFrame) -> pd.DataFrame:
    profesionales = df.copy()
    equivalencias = {
        "Documento": "CEDULA",
        "DOCUMENTO": "CEDULA",
        "Nombre": "NOMBRE",
        "Correo": "CORREO",
    }
    profesionales = profesionales.rename(columns=equivalencias)
    for columna in ("CEDULA", "NOMBRE", "CORREO"):
        if columna not in profesionales:
            profesionales[columna] = ""

    profesionales = profesionales[["CEDULA", "NOMBRE", "CORREO"]].fillna("")
    profesionales["CEDULA"] = profesionales["CEDULA"].map(limpiar_cedula)
    profesionales["NOMBRE"] = profesionales["NOMBRE"].astype(str).str.strip()
    profesionales["CORREO"] = profesionales["CORREO"].astype(str).str.strip()
    vacias = (profesionales == "").all(axis=1)
    profesionales = profesionales.loc[~vacias].reset_index(drop=True)
    return profesionales


def cargar_profesionales(ruta: str | Path) -> pd.DataFrame:
    ruta = Path(ruta)
    if not ruta.exists():
        profesionales = profesionales_por_defecto()
        guardar_profesionales(profesionales, ruta)
        return profesionales

    try:
        profesionales = pd.read_csv(
            ruta, dtype=str, keep_default_na=False, encoding="utf-8-sig"
        )
    except Exception as exc:
        raise ErrorDatos(
            f"No fue posible leer la tabla de profesionales en {ruta}. Detalle: {exc}"
        ) from exc
    return normalizar_profesionales(profesionales)


def guardar_profesionales(df: pd.DataFrame, ruta: str | Path) -> pd.DataFrame:
    ruta = Path(ruta)
    profesionales = normalizar_profesionales(df)

    sin_cedula = profesionales["CEDULA"].eq("") & profesionales["NOMBRE"].ne("")
    if sin_cedula.any():
        filas = ", ".join(str(indice + 1) for indice in profesionales.index[sin_cedula])
        raise ErrorDatos(f"Hay profesionales con nombre pero sin cédula en las filas: {filas}.")

    duplicadas = profesionales.loc[
        profesionales["CEDULA"].ne("") & profesionales["CEDULA"].duplicated(False),
        "CEDULA",
    ].unique()
    if len(duplicadas):
        raise ErrorDatos(
            "Hay cédulas duplicadas en Profesionales: " + ", ".join(duplicadas)
        )

    ruta.parent.mkdir(parents=True, exist_ok=True)
    profesionales.to_csv(ruta, index=False, encoding="utf-8-sig")
    return profesionales


def _serie_fechas(valores: pd.Series) -> pd.Series:
    try:
        return pd.to_datetime(valores, errors="coerce", format="mixed")
    except TypeError:
        return pd.to_datetime(valores, errors="coerce")


def fecha_actual_colombia() -> datetime:
    return datetime.now(ZoneInfo("America/Bogota")).replace(tzinfo=None)


def extraer_codigo_estado(valor: object) -> str:
    texto = "" if valor is None else str(valor).strip()
    if not texto or texto.casefold() == "nan":
        return ""
    coincidencia = RE_CODIGO_ESTADO.match(texto)
    return coincidencia.group(1) if coincidencia else texto


def _fechas_locales_colombia(valores: pd.Series) -> pd.Series:
    fechas = _serie_fechas(valores)
    if fechas.empty:
        return fechas
    if fechas.dt.tz is not None:
        fechas = fechas.dt.tz_convert("America/Bogota").dt.tz_localize(None)
    return fechas


def _resolver_llegada_actual(eventos: pd.DataFrame, flujo: str) -> pd.Timestamp:
    """Obtiene la llegada más reciente recorriendo devoluciones y reaperturas."""
    if eventos.empty or flujo not in REGLAS_LLEGADA:
        return pd.NaT

    regla = REGLAS_LLEGADA[flujo]
    activos = regla["activos"]
    salidas = regla["salidas"]
    ordenados = eventos.sort_values(
        ["timestamp", "_row_order"],
        kind="stable",
    )

    llegada = pd.NaT
    esperando_reingreso = False

    for evento in ordenados.itertuples(index=False):
        estado_inicial = evento.initial_state_code
        estado_final = evento.final_state_code

        if estado_final in salidas:
            esperando_reingreso = True
            continue

        if estado_final not in activos:
            continue

        es_reingreso = esperando_reingreso or estado_inicial in salidas
        if pd.isna(llegada) or es_reingreso:
            llegada = pd.Timestamp(evento.timestamp)
        esperando_reingreso = False

    return llegada


def calcular_llegadas_desde_logs(
    asignador: pd.DataFrame,
    exported_logs: pd.DataFrame,
) -> pd.Series:
    qrs = set(asignador["QR"].dropna().astype(str).str.strip())
    logs = exported_logs[
        ["request_code", "initial_state", "final_state", "date_created"]
    ].copy()
    logs["_row_order"] = range(len(logs))
    logs["request_code"] = logs["request_code"].astype(str).str.strip()
    logs = logs[logs["request_code"].isin(qrs)].copy()
    logs["timestamp"] = _fechas_locales_colombia(logs["date_created"])
    logs["initial_state_code"] = logs["initial_state"].map(extraer_codigo_estado)
    logs["final_state_code"] = logs["final_state"].map(extraer_codigo_estado)
    logs = logs.dropna(subset=["timestamp"])
    logs = logs[
        logs["initial_state_code"].isin(ESTADOS_SEGUIDOS_LLEGADA)
        | logs["final_state_code"].isin(ESTADOS_SEGUIDOS_LLEGADA)
    ].copy()
    logs = logs.sort_values(
        ["request_code", "timestamp", "_row_order"],
        kind="stable",
    )

    eventos_por_qr = {
        qr: eventos
        for qr, eventos in logs.groupby("request_code", sort=False)
    }
    llegadas: dict[tuple[str, str], pd.Timestamp] = {}
    for qr, flujo in asignador[["QR", "FLUJO"]].drop_duplicates().itertuples(
        index=False, name=None
    ):
        eventos = eventos_por_qr.get(qr)
        llegadas[(qr, flujo)] = (
            _resolver_llegada_actual(eventos, flujo)
            if eventos is not None
            else pd.NaT
        )

    valores = [
        llegadas.get((qr, flujo), pd.NaT)
        for qr, flujo in asignador[["QR", "FLUJO"]].itertuples(
            index=False, name=None
        )
    ]
    return pd.Series(pd.to_datetime(valores, errors="coerce"), index=asignador.index)


def construir_asignador(
    exported_request_dw: pd.DataFrame,
    exported_drawing: pd.DataFrame,
    grouped_node_report: pd.DataFrame,
    exported_logs: pd.DataFrame | None,
    profesionales: pd.DataFrame,
    fecha_carga: datetime | None = None,
) -> pd.DataFrame:
    validar_columnas(
        exported_request_dw,
        COLUMNAS_REQUERIDAS["exported_request_DW"],
        "exported-request_DW",
    )
    validar_columnas(
        exported_drawing,
        COLUMNAS_REQUERIDAS["exported_drawing"],
        "exported-drawing",
    )
    validar_columnas(
        grouped_node_report,
        COLUMNAS_REQUERIDAS["grouped_node_report"],
        "grouped-node-report",
    )
    if exported_logs is not None:
        validar_columnas(
            exported_logs,
            COLUMNAS_REQUERIDAS["exported_logs"],
            "exported-logs",
        )

    profesionales = normalizar_profesionales(profesionales)
    estados_viejos = exported_request_dw["status"].fillna("").astype(str).str.strip()
    estados_nuevos = exported_drawing["status"].map(normalizar_estado_flujo_nuevo)

    viejos = exported_request_dw.loc[
        exported_request_dw["status"].map(es_estado_flujo_viejo)
    ].copy()
    nuevos = exported_drawing.loc[estados_nuevos.ne("")].copy()

    asignador_viejo = pd.DataFrame(
        {
            "QR": viejos["request_code"].astype(str).str.strip(),
            "FLUJO": "Flujo viejo",
            "ORIGEN": viejos["process__subprocess"].map(normalizar_origen),
            "ESTADO": estados_viejos.loc[viejos.index],
            "USER": viejos["drawing_user_last_name"].map(
                formatear_nombre_flujo_viejo
            ),
        }
    )

    mapa_profesionales = (
        profesionales.loc[profesionales["CEDULA"].ne("")]
        .drop_duplicates("CEDULA", keep="last")
        .set_index("CEDULA")["NOMBRE"]
        .to_dict()
    )
    cedulas_nuevas = nuevos["assigned_draftsman"].map(limpiar_cedula)
    asignador_nuevo = pd.DataFrame(
        {
            "QR": nuevos["request_code"].astype(str).str.strip(),
            "FLUJO": "Flujo nuevo",
            "ORIGEN": nuevos["process.subprocess"].map(normalizar_origen),
            "ESTADO": estados_nuevos.loc[nuevos.index],
            "USER": cedulas_nuevas.map(mapa_profesionales).fillna(""),
        }
    )

    asignador = pd.concat([asignador_viejo, asignador_nuevo], ignore_index=True)
    asignador = asignador.loc[asignador["QR"].ne("")].copy()

    conteo_elementos = (
        grouped_node_report["request_code"]
        .astype(str)
        .str.strip()
        .value_counts()
    )
    asignador["ELEMENTOS"] = (
        asignador["QR"].map(conteo_elementos).fillna(0).astype("int64")
    )

    if exported_logs is None:
        asignador["LLEGADA"] = pd.Timestamp(fecha_carga or fecha_actual_colombia())
    else:
        asignador["LLEGADA"] = calcular_llegadas_desde_logs(
            asignador,
            exported_logs,
        )

    return asignador[COLUMNAS_ASIGNADOR].reset_index(drop=True)


def construir_resumen(asignador: pd.DataFrame) -> dict[str, pd.DataFrame]:
    datos = asignador.copy()
    datos["USER_RESUMEN"] = datos["USER"].fillna("").replace("", "Sin asignar")

    casos_flujo = (
        datos.groupby(["FLUJO", "ESTADO"], dropna=False)["QR"]
        .count()
        .reset_index(name="CASOS")
        .sort_values(["FLUJO", "ESTADO"])
        .reset_index(drop=True)
    )
    casos_origen = (
        datos.groupby(["ORIGEN", "ESTADO"], dropna=False)["QR"]
        .count()
        .reset_index(name="CASOS")
        .sort_values(["ORIGEN", "ESTADO"])
        .reset_index(drop=True)
    )
    casos_user = (
        datos.groupby(["USER_RESUMEN", "ESTADO"], dropna=False)["QR"]
        .count()
        .reset_index(name="CASOS")
        .rename(columns={"USER_RESUMEN": "USER"})
        .sort_values(["USER", "ESTADO"])
        .reset_index(drop=True)
    )
    elementos_user = (
        datos.groupby(["USER_RESUMEN", "ESTADO"], dropna=False)["ELEMENTOS"]
        .sum()
        .reset_index(name="ELEMENTOS")
        .rename(columns={"USER_RESUMEN": "USER"})
        .sort_values(["USER", "ESTADO"])
        .reset_index(drop=True)
    )
    return {
        "casos_por_flujo": casos_flujo,
        "casos_por_origen": casos_origen,
        "casos_por_user": casos_user,
        "elementos_por_user": elementos_user,
    }


def _ancho_columnas(df: pd.DataFrame, minimo: int = 10, maximo: int = 42) -> list[int]:
    anchos = []
    muestra = df.head(2_000)
    for columna in df.columns:
        largos = muestra[columna].map(
            lambda valor: 0 if valor is None or pd.isna(valor) else len(str(valor))
        )
        largo = max(
            len(str(columna)),
            int(largos.max()) if not largos.empty else 0,
        )
        anchos.append(max(minimo, min(maximo, int(largo) + 2)))
    return anchos


COLOR_AZUL = "1F4E78"
COLOR_BORDE = "B4C6E7"
RE_CARACTERES_ILEGALES = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def _valor_excel(valor: object) -> object:
    if valor is None or pd.isna(valor):
        return None
    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime().replace(tzinfo=None)
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    if hasattr(valor, "item") and not isinstance(valor, str):
        try:
            valor = valor.item()
        except ValueError:
            pass
    if isinstance(valor, str):
        return RE_CARACTERES_ILEGALES.sub("", valor)[:32_767]
    return valor


def _celda(
    hoja: object,
    valor: object,
    *,
    encabezado: bool = False,
    titulo: bool = False,
    numero: bool = False,
    fecha: bool = False,
) -> WriteOnlyCell:
    celda = WriteOnlyCell(hoja, value=_valor_excel(valor))
    if isinstance(celda.value, str):
        celda.data_type = "s"
    if encabezado or titulo:
        celda.font = Font(bold=True, color="FFFFFF", size=12 if titulo else 11)
        celda.fill = PatternFill("solid", fgColor=COLOR_AZUL)
        celda.alignment = Alignment(
            horizontal="left" if titulo else "center",
            vertical="center",
        )
        celda.border = Border(
            left=Side(style="thin", color=COLOR_BORDE),
            right=Side(style="thin", color=COLOR_BORDE),
            top=Side(style="thin", color=COLOR_BORDE),
            bottom=Side(style="thin", color=COLOR_BORDE),
        )
    if numero:
        celda.number_format = "0"
    if fecha:
        celda.number_format = "yyyy-mm-dd hh:mm:ss"
    return celda


def _configurar_hoja(hoja: object, congelar: str = "A2") -> None:
    hoja.freeze_panes = congelar
    hoja.sheet_view.showGridLines = False
    hoja.page_setup.orientation = "landscape"
    hoja.oddHeader.center.text = "Casos para asignación"
    hoja.oddHeader.center.size = 10
    hoja.oddHeader.center.font = "Arial,Bold"
    hoja.oddFooter.left.text = "Generado por Asignador"
    hoja.oddFooter.center.text = "Página &P de &N"


def _escribir_dataframe(
    hoja: object,
    df: pd.DataFrame,
    *,
    incluir_encabezado: bool = True,
    columnas_fecha: set[str] | None = None,
    columnas_numero: set[str] | None = None,
) -> None:
    columnas_fecha = columnas_fecha or set()
    columnas_numero = columnas_numero or set()
    columnas = list(df.columns)
    if incluir_encabezado:
        hoja.append([_celda(hoja, columna, encabezado=True) for columna in columnas])
    for fila in df.itertuples(index=False, name=None):
        hoja.append(
            [
                _celda(
                    hoja,
                    valor,
                    fecha=columna in columnas_fecha,
                    numero=columna in columnas_numero,
                )
                for columna, valor in zip(columnas, fila)
            ]
        )


def _ajustar_anchos(
    hoja: object, df: pd.DataFrame, minimo: int = 10, maximo: int = 42
) -> None:
    for indice, ancho in enumerate(
        _ancho_columnas(df, minimo=minimo, maximo=maximo), start=1
    ):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho


def _escribir_raw(
    workbook: Workbook,
    nombre_base: str,
    df: pd.DataFrame,
) -> list[str]:
    hojas = []
    total = max(1, (len(df) + MAX_FILAS_DATOS_EXCEL - 1) // MAX_FILAS_DATOS_EXCEL)
    for numero in range(total):
        inicio = numero * MAX_FILAS_DATOS_EXCEL
        fin = min(len(df), (numero + 1) * MAX_FILAS_DATOS_EXCEL)
        nombre = nombre_base if numero == 0 else f"{nombre_base}_{numero + 1}"
        nombre = nombre[:31]
        fragmento = df.iloc[inicio:fin]
        hoja = workbook.create_sheet(nombre)
        _configurar_hoja(hoja)
        _escribir_dataframe(hoja, fragmento)
        ultima_columna = get_column_letter(max(1, len(fragmento.columns)))
        hoja.auto_filter.ref = f"A1:{ultima_columna}{len(fragmento) + 1}"
        hoja.row_dimensions[1].height = 24
        _ajustar_anchos(hoja, fragmento, maximo=34)
        hojas.append(nombre)
    return hojas


def exportar_excel(
    asignador: pd.DataFrame,
    resumenes: Mapping[str, pd.DataFrame],
    profesionales: pd.DataFrame,
    raw_data: Mapping[str, pd.DataFrame],
) -> bytes:
    salida = io.BytesIO()
    workbook = Workbook(write_only=True, iso_dates=True)

    hoja_asignador = workbook.create_sheet("Asignador")
    _configurar_hoja(hoja_asignador)
    _escribir_dataframe(
        hoja_asignador,
        asignador,
        columnas_fecha={"LLEGADA"},
        columnas_numero={"ELEMENTOS"},
    )
    hoja_asignador.auto_filter.ref = (
        f"A1:{get_column_letter(len(asignador.columns))}{len(asignador) + 1}"
    )
    hoja_asignador.row_dimensions[1].height = 24
    _ajustar_anchos(hoja_asignador, asignador)

    hoja_resumen = workbook.create_sheet("Resumen")
    _configurar_hoja(hoja_resumen, congelar="A1")
    hoja_resumen.sheet_properties.tabColor = COLOR_AZUL
    tablas = [
        ("Casos por flujo", resumenes["casos_por_flujo"]),
        ("Casos por origen", resumenes["casos_por_origen"]),
        ("Casos por USER", resumenes["casos_por_user"]),
        ("Elementos por USER", resumenes["elementos_por_user"]),
    ]
    for titulo, tabla in tablas:
        hoja_resumen.append(
            [_celda(hoja_resumen, titulo, titulo=True)]
            + [_celda(hoja_resumen, "", titulo=True) for _ in range(2)]
        )
        _escribir_dataframe(
            hoja_resumen,
            tabla,
            columnas_numero={tabla.columns[-1]} if len(tabla.columns) else set(),
        )
        hoja_resumen.append([])
        hoja_resumen.append([])
    hoja_resumen.column_dimensions["A"].width = 38
    hoja_resumen.column_dimensions["B"].width = 24
    hoja_resumen.column_dimensions["C"].width = 14

    profesionales = normalizar_profesionales(profesionales)
    hoja_profesionales = workbook.create_sheet("Profesionales")
    _configurar_hoja(hoja_profesionales)
    _escribir_dataframe(hoja_profesionales, profesionales)
    hoja_profesionales.auto_filter.ref = (
        f"A1:C{len(profesionales) + 1}"
    )
    hoja_profesionales.row_dimensions[1].height = 24
    hoja_profesionales.column_dimensions["A"].width = 18
    hoja_profesionales.column_dimensions["B"].width = 42
    hoja_profesionales.column_dimensions["C"].width = 38

    nombres_raw = {
        "exported_logs": "Raw_exported_logs",
        "grouped_node_report": "Raw_grouped_node_report",
        "exported_request_DW": "Raw_exported_request_DW",
        "exported_drawing": "Raw_exported_drawing",
    }
    for tipo, nombre_hoja in nombres_raw.items():
        if tipo not in raw_data:
            raise ErrorDatos(f"Faltan datos Raw para {tipo}.")
        _escribir_raw(workbook, nombre_hoja, raw_data[tipo])

    workbook.save(salida)
    return salida.getvalue()


def procesar_archivos(
    archivos_identificados: Mapping[str, object],
    profesionales: pd.DataFrame,
) -> tuple[
    pd.DataFrame,
    dict[str, pd.DataFrame],
    dict[str, pd.DataFrame],
]:
    faltantes = [
        tipo for tipo in ARCHIVOS_OBLIGATORIOS if tipo not in archivos_identificados
    ]
    if faltantes:
        etiquetas = [TIPOS_ARCHIVO[tipo] for tipo in faltantes]
        raise ErrorDatos("Faltan archivos requeridos: " + ", ".join(etiquetas))

    raw_data = {
        tipo: leer_csv(archivo)
        for tipo, archivo in archivos_identificados.items()
        if tipo in TIPOS_ARCHIVO
    }
    for tipo, df in raw_data.items():
        validar_columnas(
            df,
            COLUMNAS_REQUERIDAS[tipo],
            _nombre_archivo(archivos_identificados[tipo]),
        )

    asignador = construir_asignador(
        raw_data["exported_request_DW"],
        raw_data["exported_drawing"],
        raw_data["grouped_node_report"],
        raw_data.get("exported_logs"),
        profesionales,
        fecha_carga=fecha_actual_colombia(),
    )
    resumenes = construir_resumen(asignador)
    raw_data.setdefault(
        "exported_logs",
        pd.DataFrame(
            columns=[
                "request_code",
                "initial_state",
                "final_state",
                "date_created",
            ]
        ),
    )
    return asignador, resumenes, raw_data
